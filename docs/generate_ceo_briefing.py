"""Generate CEO briefing PPTX + PDF for Workforce platform."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from reportlab.lib.colors import Color, HexColor, white, black
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch, mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
    HRFlowable,
    ListFlowable,
    ListItem,
)
from reportlab.lib.enums import TA_LEFT, TA_CENTER

OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "Workforce_CEO_Briefing.pptx"
PDF_PATH = OUT_DIR / "Workforce_CEO_Briefing.pdf"

# Brand palette — professional teal/slate (not purple AI default)
NAVY = HexColor("#0F2A3D")
TEAL = HexColor("#1A6B6B")
ACCENT = HexColor("#C45C26")
LIGHT = HexColor("#F4F7F8")
MUTED = HexColor("#5A6B75")
LINE = HexColor("#D5DEE3")

PPTX_NAVY = RGBColor(0x0F, 0x2A, 0x3D)
PPTX_TEAL = RGBColor(0x1A, 0x6B, 0x6B)
PPTX_ACCENT = RGBColor(0xC4, 0x5C, 0x26)
PPTX_LIGHT = RGBColor(0xF4, 0xF7, 0xF8)
PPTX_MUTED = RGBColor(0x5A, 0x6B, 0x75)
PPTX_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_run(run, size=18, bold=False, color=PPTX_NAVY, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_bg(slide, color=PPTX_LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_top_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = PPTX_TEAL
    bar.line.fill.background()


def _add_footer(slide, page: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(10), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"Workforce  ·  Confidential CEO Briefing  ·  {page}/{total}"
    _set_run(run, size=10, color=PPTX_MUTED)


def _title_block(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run(run, size=28, bold=True, color=PPTX_NAVY)
    if subtitle:
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = subtitle
        _set_run(run2, size=14, color=PPTX_MUTED)


def _bullets(slide, items: list[str], left=0.6, top=1.4, width=12, height=5.2, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {item}"
        _set_run(run, size=size, color=PPTX_NAVY)


def _two_col_cards(slide, left_title, left_items, right_title, right_items, top=1.35):
    for idx, (title, items) in enumerate(
        [(left_title, left_items), (right_title, right_items)]
    ):
        x = 0.55 if idx == 0 else 6.95
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(5.9), Inches(5.2)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = PPTX_WHITE
        card.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE3)

        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(top), Inches(0.12), Inches(5.2)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = PPTX_TEAL if idx == 0 else PPTX_ACCENT
        accent.line.fill.background()

        tbox = slide.shapes.add_textbox(
            Inches(x + 0.35), Inches(top + 0.25), Inches(5.3), Inches(0.5)
        )
        tr = tbox.text_frame.paragraphs[0].add_run()
        tr.text = title
        _set_run(tr, size=18, bold=True, color=PPTX_NAVY)

        ibox = slide.shapes.add_textbox(
            Inches(x + 0.35), Inches(top + 0.9), Inches(5.3), Inches(4.0)
        )
        tf = ibox.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            run = p.add_run()
            run.text = f"•  {item}"
            _set_run(run, size=13, color=PPTX_NAVY)


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides_meta = []

    # 1 Title
    s = prs.slides.add_slide(blank)
    _add_bg(s, PPTX_NAVY)
    bar = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = PPTX_TEAL
    bar.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11), Inches(3.5))
    tf = box.text_frame
    r = tf.paragraphs[0].add_run()
    r.text = "WORKFORCE"
    _set_run(r, size=18, bold=True, color=PPTX_TEAL)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Office Attendance & People Operations"
    _set_run(r2, size=36, bold=True, color=PPTX_WHITE)
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = (
        "CEO briefing: how the system works, what it solves,\n"
        "and the foundation for long-term growth"
    )
    _set_run(r3, size=16, color=RGBColor(0xB8, 0xC5, 0xCD))
    p4 = tf.add_paragraph()
    p4.space_before = Pt(24)
    r4 = p4.add_run()
    r4.text = "Simple to understand  ·  Strong operational foundation"
    _set_run(r4, size=13, color=PPTX_TEAL)
    slides_meta.append(s)

    # 2 One-sentence purpose
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "What this project is", "One clear purpose")
    card = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.0), Inches(12), Inches(3.2)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = PPTX_WHITE
    card.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE3)
    t = s.shapes.add_textbox(Inches(1.1), Inches(2.5), Inches(11.2), Inches(2.2))
    tf = t.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = (
        "Workforce proves people are at the office when they should be, "
        "helps managers see who is present, late, or missing checkout, "
        "and connects the same company people into chat, daily tasks, "
        "and performance — without spreadsheets and guesswork."
    )
    _set_run(r, size=20, color=PPTX_NAVY)
    slides_meta.append(s)

    # 3 Problems solved
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "Main problems we solve", "Why the business needs this")
    _bullets(
        s,
        [
            "No reliable proof that staff are physically at the office",
            "Managers waste time asking “who is in / late / forgot to leave?”",
            "Leave, holidays, and attendance conflict or live in separate tools",
            "Task work and performance sit outside the same people record",
            "Lobby boards and reports require manual updates",
            "Growing companies need one system of truth per office and company",
        ],
    )
    slides_meta.append(s)

    # 4 Product surfaces
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "Four surfaces — one system", "Each role gets the right tool")
    cards = [
        ("Workforce Control", "Admins configure offices,\nschedules, staff, leave,\nreports & performance"),
        ("Work-Force App", "Employees check in/out\nwith location, request\nleave, chat, see reviews"),
        ("Office Display", "Lobby tablets show\nwho is present today\nand meeting rooms"),
        ("Task Operations", "Same company people\nrun projects & boards\nwithout a second login world"),
    ]
    for i, (title, body) in enumerate(cards):
        x = 0.45 + i * 3.2
        shape = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.6), Inches(3.0), Inches(4.6)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = PPTX_WHITE
        shape.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE3)
        top = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.6), Inches(3.0), Inches(0.12)
        )
        top.fill.solid()
        top.fill.fore_color.rgb = PPTX_TEAL
        top.line.fill.background()
        tb = s.shapes.add_textbox(Inches(x + 0.2), Inches(2.0), Inches(2.6), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        rt = tf.paragraphs[0].add_run()
        rt.text = title
        _set_run(rt, size=16, bold=True, color=PPTX_NAVY)
        p = tf.add_paragraph()
        p.space_before = Pt(14)
        rb = p.add_run()
        rb.text = body
        _set_run(rb, size=13, color=PPTX_MUTED)
    slides_meta.append(s)

    # 5 Office setup
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(
        s,
        "Office setup — the foundation",
        "Correct setup once → every day works automatically",
    )
    _bullets(
        s,
        [
            "Create company (organization) → add offices on a map with allowed GPS radius",
            "Define work schedules (check-in / check-out times per weekday)",
            "Add departments and employees; each person is linked to one office + schedule",
            "Set company rules: photo required, auto-checkout, desktop location options",
            "Assign office admins for day-to-day control of their site(s)",
            "Invite staff — they join the right office with the right rules from day one",
        ],
        size=15,
    )
    note = s.shapes.add_textbox(Inches(0.6), Inches(6.2), Inches(12), Inches(0.5))
    nr = note.text_frame.paragraphs[0].add_run()
    nr.text = "Key idea: Office + Schedule = the rules of attendance. Everything else builds on this."
    _set_run(nr, size=13, bold=True, color=PPTX_ACCENT)
    slides_meta.append(s)

    # 6 Check-in / out workflow
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "Daily check-in / check-out", "How attendance actually works")
    steps = [
        ("1", "Preview", "Employee opens app\nand confirms location"),
        ("2", "Check in", "GPS must be inside\noffice radius (+ photo)"),
        ("3", "Work day", "Late reasons, reminders,\npresence on admin & display"),
        ("4", "Check out", "Worked time recorded;\noptional daily worksheet"),
        ("5", "Safety net", "Auto / missing checkout\njobs close open days"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.4 + i * 2.55
        circ = s.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.85), Inches(1.55), Inches(0.55), Inches(0.55)
        )
        circ.fill.solid()
        circ.fill.fore_color.rgb = PPTX_TEAL
        circ.line.fill.background()
        nb = s.shapes.add_textbox(Inches(x + 0.85), Inches(1.62), Inches(0.55), Inches(0.45))
        np = nb.text_frame.paragraphs[0]
        np.alignment = PP_ALIGN.CENTER
        nr = np.add_run()
        nr.text = num
        _set_run(nr, size=14, bold=True, color=PPTX_WHITE)
        card = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.35), Inches(2.4), Inches(3.5)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = PPTX_WHITE
        card.line.color.rgb = RGBColor(0xD5, 0xDE, 0xE3)
        tb = s.shapes.add_textbox(Inches(x + 0.15), Inches(2.6), Inches(2.1), Inches(3.0))
        tf = tb.text_frame
        tf.word_wrap = True
        rt = tf.paragraphs[0].add_run()
        rt.text = title
        _set_run(rt, size=15, bold=True, color=PPTX_NAVY)
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        rb = p.add_run()
        rb.text = body
        _set_run(rb, size=12, color=PPTX_MUTED)
    slides_meta.append(s)

    # 7 Attendance rules
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "Attendance rules that keep data clean", "Trustworthy records, not vague timestamps")
    _two_col_cards(
        s,
        "What the system enforces",
        [
            "Must be inside the office geofence",
            "Blocked on approved leave or holiday",
            "Server time wins (not device clock)",
            "Late reason on first late entry",
            "Re-entry only after proper checkout",
        ],
        "What managers gain",
        [
            "Who is present right now",
            "Late / early / overtime minutes",
            "Missing checkout alerts",
            "Correction / correctness requests",
            "Realtime updates to admin & display",
        ],
    )
    slides_meta.append(s)

    # 8 Chat
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "Chat — operational messaging", "Same company people, right scope")
    _bullets(
        s,
        [
            "Direct messages, group chats, and admin ↔ employee channels",
            "Scoped by company; office admins see colleagues for their office(s)",
            "Not a timesheet chatbot — messaging sits beside attendance, not instead of it",
            "Keeps day-to-day coordination inside Workforce (less WhatsApp sprawl for ops)",
            "Same directory as attendance, tasks, and performance — one identity",
        ],
        size=15,
    )
    slides_meta.append(s)

    # 9 Task operations
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(
        s,
        "Task Operations — how daily work runs",
        "Projects and boards using the same HR people list",
    )
    _two_col_cards(
        s,
        "What it is",
        [
            "Company task board / Kanban-style work",
            "Projects, schedules, and assignments",
            "Staff synced from Workforce employees",
            "Opened from admin or employee app",
            "Own permission levels for task work",
        ],
        "Why it matters",
        [
            "Attendance proves presence; tasks prove output",
            "No second people database to maintain",
            "Managers see people + work in one family of tools",
            "Ready to grow without rebuilding identity",
            "Clear separation: HR rules vs task roles",
        ],
    )
    slides_meta.append(s)

    # 10 Performance
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(
        s,
        "Performance — fairness with evidence",
        "Reviews that can use attendance reliability",
    )
    _bullets(
        s,
        [
            "Admins run evaluation cycles with templates, scores, and goals",
            "System can score reliability from attendance facts (lates, missing checkouts, worksheets, leave)",
            "Employees can self-score where the company allows it",
            "Connects “were you present and dependable?” to formal performance — not gut feel only",
            "Same org people record — performance stays linked to the real employee",
        ],
        size=15,
    )
    slides_meta.append(s)

    # 11 End-to-end flow
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "End-to-end workflow", "From setup to long-term value")
    flow = [
        "1. Admin sets office + schedule + staff",
        "2. Employee checks in / out each day",
        "3. Leave & holidays protect the calendar",
        "4. Chat keeps the office talking",
        "5. Task Operations runs the work",
        "6. Performance reviews use real signals",
        "7. Display board shows who’s in",
        "8. Reports & corrections keep trust",
    ]
    _bullets(s, flow, size=15)
    slides_meta.append(s)

    # 12 Assumptions
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "Assumptions we build on", "Honest foundations for the CEO")
    _bullets(
        s,
        [
            "Multi-company (multi-tenant): each organization is isolated",
            "Physical presence matters — GPS geofence is the proof of being at the office",
            "One primary office + work schedule per employee is the default model",
            "Leave and holidays should block check-in (no double counting the day)",
            "Ethiopia-first defaults (timezone, holiday calendar, bilingual employee UX) — expandable",
            "Tasks are operational work; they do not replace attendance proof",
            "Roles are clear: Platform → Company Admin → Office Admin → Employee",
        ],
        size=14,
    )
    slides_meta.append(s)

    # 13 Long-term
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(
        s,
        "Long-term direction",
        "Not complex today — designed to grow without rebuild",
    )
    _two_col_cards(
        s,
        "Strong foundation already",
        [
            "One backend for all apps",
            "Permission codes + office scope",
            "Feature flags per company",
            "Realtime (admin + display)",
            "Audit trail & correction flows",
            "Embedded Task Operations",
        ],
        "Natural next horizon",
        [
            "More offices / tenants at scale",
            "Deeper analytics for leaders",
            "Richer performance & goals",
            "Stronger display / meeting ops",
            "More automation (policies)",
            "Modules without a second identity",
        ],
    )
    slides_meta.append(s)

    # 14 CEO takeaway
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _add_top_bar(s)
    _title_block(s, "CEO takeaway", "Easy to explain. Hard to fake. Ready to scale.")
    _bullets(
        s,
        [
            "Workforce is the system of truth for office presence and people operations",
            "Setup is simple: map the office, set hours, assign people — then the day runs",
            "Check-in/out + leave + chat + tasks + performance share one company identity",
            "Problem solved: less guesswork, less spreadsheet ops, more trustworthy records",
            "Foundation is strong: clear roles, geofence proof, automation safety nets, room to grow",
        ],
        size=15,
    )
    slides_meta.append(s)

    # 15 Closing
    s = prs.slides.add_slide(blank)
    _add_bg(s, PPTX_NAVY)
    bar = s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = PPTX_TEAL
    bar.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11), Inches(2.5))
    tf = box.text_frame
    r = tf.paragraphs[0].add_run()
    r.text = "Ready for discussion"
    _set_run(r, size=32, bold=True, color=PPTX_WHITE)
    p2 = tf.add_paragraph()
    p2.space_before = Pt(12)
    r2 = p2.add_run()
    r2.text = "Workforce · Office attendance & people operations\nQuestions, priorities, and next milestones welcome."
    _set_run(r2, size=15, color=RGBColor(0xB8, 0xC5, 0xCD))
    slides_meta.append(s)

    total = len(slides_meta)
    for i, slide in enumerate(slides_meta, start=1):
        # Skip dark title/closing decorative footer contrast
        if i in (1, total):
            continue
        _add_footer(slide, i, total)

    prs.save(str(PPTX_PATH))
    return total


def build_pdf():
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=landscape(A4),
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = {
        "h1": ParagraphStyle(
            "h1",
            fontName="Helvetica-Bold",
            fontSize=22,
            textColor=NAVY,
            spaceAfter=6,
        ),
        "sub": ParagraphStyle(
            "sub",
            fontName="Helvetica",
            fontSize=11,
            textColor=MUTED,
            spaceAfter=14,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Helvetica",
            fontSize=11,
            textColor=NAVY,
            leading=16,
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "bullet",
            fontName="Helvetica",
            fontSize=11,
            textColor=NAVY,
            leading=15,
            leftIndent=12,
            spaceAfter=4,
        ),
        "cover": ParagraphStyle(
            "cover",
            fontName="Helvetica-Bold",
            fontSize=28,
            textColor=white,
            leading=34,
        ),
        "coverSub": ParagraphStyle(
            "coverSub",
            fontName="Helvetica",
            fontSize=12,
            textColor=HexColor("#B8C5CD"),
            leading=18,
        ),
        "label": ParagraphStyle(
            "label",
            fontName="Helvetica-Bold",
            fontSize=10,
            textColor=TEAL,
            spaceAfter=8,
        ),
        "cardTitle": ParagraphStyle(
            "cardTitle",
            fontName="Helvetica-Bold",
            fontSize=12,
            textColor=NAVY,
            spaceAfter=6,
        ),
    }

    story = []

    def section(title, subtitle, bullets):
        block = [
            Paragraph(title, styles["h1"]),
            Paragraph(subtitle, styles["sub"]),
            HRFlowable(width="100%", thickness=1, color=LINE, spaceAfter=10),
        ]
        for b in bullets:
            block.append(Paragraph(f"•  {b}", styles["bullet"]))
        block.append(Spacer(1, 18))
        story.append(KeepTogether(block))

    # Cover
    cover_data = [
        [
            Paragraph("WORKFORCE", styles["label"]),
        ],
        [
            Paragraph(
                "Office Attendance & People Operations",
                styles["cover"],
            )
        ],
        [
            Paragraph(
                "CEO briefing: how the system works, what it solves,<br/>and the foundation for long-term growth",
                styles["coverSub"],
            )
        ],
        [
            Paragraph(
                "Simple to understand  ·  Strong operational foundation",
                styles["label"],
            )
        ],
    ]
    cover = Table(cover_data, colWidths=[9.5 * inch])
    cover.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), NAVY),
                ("TOPPADDING", (0, 0), (-1, -1), 14),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
                ("LEFTPADDING", (0, 0), (-1, -1), 24),
                ("RIGHTPADDING", (0, 0), (-1, -1), 24),
            ]
        )
    )
    story.append(cover)
    story.append(Spacer(1, 22))

    story.append(Paragraph("What this project is", styles["h1"]))
    story.append(Paragraph("One clear purpose", styles["sub"]))
    story.append(
        Paragraph(
            "Workforce proves people are at the office when they should be, helps managers see "
            "who is present, late, or missing checkout, and connects the same company people into "
            "chat, daily tasks, and performance — without spreadsheets and guesswork.",
            styles["body"],
        )
    )
    story.append(Spacer(1, 16))

    section(
        "Main problems we solve",
        "Why the business needs this",
        [
            "No reliable proof that staff are physically at the office",
            "Managers waste time asking who is in, late, or forgot to leave",
            "Leave, holidays, and attendance conflict or live in separate tools",
            "Task work and performance sit outside the same people record",
            "Lobby boards and reports require manual updates",
            "Growing companies need one system of truth per office and company",
        ],
    )

    section(
        "Four surfaces — one system",
        "Each role gets the right tool",
        [
            "Workforce Control (admin): configure offices, schedules, staff, leave, reports, performance",
            "Work-Force app (employee): geo check-in/out, leave, chat, evaluations",
            "Office Display: lobby tablets show who is present and meeting rooms",
            "Task Operations: same company people run projects and boards",
        ],
    )

    section(
        "Office setup — the foundation",
        "Correct setup once → every day works automatically",
        [
            "Create company → add offices on a map with allowed GPS radius",
            "Define work schedules (check-in / check-out times per weekday)",
            "Add departments and employees; each person linked to one office + schedule",
            "Set company rules: photo required, auto-checkout, desktop location options",
            "Assign office admins; invite staff into the right office with the right rules",
            "Key idea: Office + Schedule = the rules of attendance. Everything else builds on this.",
        ],
    )

    section(
        "Daily check-in / check-out",
        "How attendance actually works",
        [
            "1 Preview location → 2 Check in inside geofence (photo if required) → 3 Work day with reminders",
            "4 Check out with worked time and optional worksheet → 5 Auto / missing-checkout jobs close open days",
            "Enforced: geofence, leave/holiday block, server time, late reason, clean re-entry rules",
            "Managers get: presence, late/early/overtime, missing checkout alerts, realtime admin + display",
        ],
    )

    section(
        "Chat — operational messaging",
        "Same company people, right scope",
        [
            "Direct, group, and admin ↔ employee channels",
            "Scoped by company; office admins see their office colleagues",
            "Messaging beside attendance — not a timesheet chatbot",
            "One identity shared with attendance, tasks, and performance",
        ],
    )

    section(
        "Task Operations",
        "How daily work runs on the same people list",
        [
            "Projects, Kanban-style boards, schedules, and assignments",
            "Staff synced from Workforce employees — no second people database",
            "Opened from admin or employee app; own permission levels for task work",
            "Attendance proves presence; tasks prove output — together they tell the full story",
        ],
    )

    section(
        "Performance",
        "Fairness with evidence",
        [
            "Evaluation cycles with templates, scores, and goals",
            "System can score reliability from attendance facts (lates, missing checkouts, worksheets, leave)",
            "Connects dependability to formal review — not gut feel only",
            "Same org people record keeps performance linked to the real employee",
        ],
    )

    section(
        "End-to-end workflow",
        "From setup to long-term value",
        [
            "Admin sets office + schedule + staff",
            "Employee checks in / out each day",
            "Leave & holidays protect the calendar",
            "Chat keeps the office talking",
            "Task Operations runs the work",
            "Performance reviews use real signals",
            "Display board shows who’s in; reports & corrections keep trust",
        ],
    )

    section(
        "Assumptions we build on",
        "Honest foundations for leadership",
        [
            "Multi-company (multi-tenant): each organization is isolated",
            "Physical presence matters — GPS geofence is the proof",
            "One primary office + work schedule per employee (default model)",
            "Leave and holidays block check-in (no double counting)",
            "Ethiopia-first defaults (timezone, holidays, bilingual UX) — expandable",
            "Tasks are operational; they do not replace attendance proof",
            "Clear roles: Platform → Company Admin → Office Admin → Employee",
        ],
    )

    section(
        "Long-term direction",
        "Not complex today — designed to grow without rebuild",
        [
            "Foundation: one backend, permission codes, office scope, feature flags, realtime, audit, Task Operations",
            "Horizon: more tenants/offices, deeper analytics, richer performance, stronger display/meeting ops, more policy automation",
            "Modules can grow without inventing a second identity system",
        ],
    )

    section(
        "CEO takeaway",
        "Easy to explain. Hard to fake. Ready to scale.",
        [
            "Workforce is the system of truth for office presence and people operations",
            "Setup is simple: map the office, set hours, assign people — then the day runs",
            "Check-in/out + leave + chat + tasks + performance share one company identity",
            "Problem solved: less guesswork, less spreadsheet ops, more trustworthy records",
            "Foundation is strong: clear roles, geofence proof, automation safety nets, room to grow",
        ],
    )

    doc.build(story)


if __name__ == "__main__":
    n = build_pptx()
    build_pdf()
    print(f"Wrote {PPTX_PATH} ({n} slides)")
    print(f"Wrote {PDF_PATH}")
